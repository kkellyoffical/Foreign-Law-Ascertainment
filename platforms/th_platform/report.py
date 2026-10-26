"""查明意见书导出：把会话的问答 + 引用 citations 组装为规范的《域外法查明意见书》。

GET /api/conversations/{conv_id}/report?format=pdf|docx
- docx 用 python-docx 构建（单一真相源）；pdf 由 LibreOffice headless 转换（zafl 可运行，
  与软著手册同一条转换链）。
- 文书结构：标题/编号/日期 → 一、查明事项 → 二、查明结论（逐问逐答）→ 三、成文法依据
  → 四、判例支撑（有材料时）→ 五、官方公告/监管指南/条约 → 六、查明方法与局限 → 附录·来源清单 → 落款。
- 引用按类型去重（statute 按 section_id、case 按 document_id、官方材料按 chunk_id），
  保持首次出现顺序。
"""
from __future__ import annotations

import logging
import os
import re
import subprocess
import tempfile
import threading
from datetime import date
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Response
from za_common import inline_citations

from . import config, db, deps

logger = logging.getLogger("my_platform.report")

router = APIRouter(prefix="/api/conversations", tags=["report"])

# 导出并发闸门：soffice 转换较重，最多同时 2 个导出，超额 429（不排队防堆积）。
_export_slots = threading.Semaphore(2)
# LibreOffice 单实例串行锁（共享 profile 并发启动易损坏）。
_soffice_lock = threading.Lock()

SOFFICE = "/usr/bin/soffice"


def _is_thai_original(source: dict) -> bool:
    return str(source.get("language") or "").lower().startswith("th")


def _statute_locator(source: dict) -> str:
    prefix = "มาตรา" if _is_thai_original(source) else "Section"
    return f"{prefix} {source.get('section_label') or ''}".strip()


def _statute_text_attribute(source: dict) -> str:
    if _is_thai_original(source):
        return "OCS 官方泰文原文"
    if source.get("translation_official") is True:
        return "官方机构提供或托管的英译"
    if source.get("translation_official") is False:
        return "非官方英译，仅供研究参考"
    return "英文译本"


# ---------------- docx 构建 ----------------

SERIF = "Noto Serif CJK SC"
MONO = "DejaVu Sans Mono"


def _cjk(run, ea: str = SERIF) -> None:
    from docx.oxml.ns import qn
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ea)


def _inline(p, text: str, size=None, bold=False, color=None) -> None:
    """极简行内 markdown，并把共享引用标记渲染为论文式上标。"""
    from docx.shared import Pt, RGBColor
    for segment in re.split(r"(\[\[CITE:\d+\]\])", text):
        marker = re.fullmatch(r"\[\[CITE:(\d+)\]\]", segment or "")
        if marker:
            r = p.add_run(marker.group(1)); r.font.superscript = True
            r.font.size = Pt(max(7.5, float(size or 10.5) * 0.74))
            r.font.color.rgb = RGBColor(0x7E, 0x18, 0x1E); _cjk(r)
            continue
        for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", segment):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                r = p.add_run(part[1:-1]); r.font.name = MONO
            else:
                r = p.add_run(part.replace("*", ""))
                r.bold = r.bold or bold
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
            _cjk(r)


def _md_to_docx(doc, md: str, anchors: list[dict] | None = None) -> None:
    """把助手回答（markdown）渲染进 docx：标题/列表/表格/引用块/分隔线。"""
    from docx.shared import Pt, RGBColor

    md = inline_citations.annotate_markdown(md or "", anchors or [])
    lines = (md or "").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        # markdown 表格块
        if line.lstrip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^[\s:\-]+$", "".join(cells)):
                    rows.append(cells)
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.paragraphs[0].text = ""
                        _inline(cell.paragraphs[0], row[ci] if ci < len(row) else "",
                                size=9, bold=(ri == 0))
            continue
        i += 1
        if not line.strip():
            continue
        if re.match(r"^\s*(---+|\*\*\*+)\s*$", line):
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            _inline(p, m.group(2).strip(), size=12 if level <= 2 else 11, bold=True)
            continue
        m = re.match(r"^\s*(?:[-*•]|\d+[.、])\s+(.*)$", line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.space_after = Pt(2)
            bullet = "· " if re.match(r"^\s*[-*•]", line) else re.match(r"^\s*(\d+[.、])", line).group(1) + " "
            r = p.add_run(bullet); r.font.size = Pt(10.5); _cjk(r)
            _inline(p, m.group(1), size=10.5)
            continue
        if line.lstrip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _inline(p, line.lstrip("> "), size=10, color=RGBColor(0x60, 0x64, 0x70))
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35
        _inline(p, line, size=10.5)


def _add_inline_notes(doc, anchors: list[dict]) -> None:
    """紧随本轮结论输出注释；文末完整来源清单仍作为审计目录保留。"""
    from docx.shared import Pt, RGBColor

    notes = inline_citations.unique_inline_notes(anchors)
    if not notes:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("随文注释")
    r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x7E, 0x18, 0x1E); _cjk(r)
    for note in notes:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(14)
        p.paragraph_format.first_line_indent = Pt(-14)
        p.paragraph_format.space_after = Pt(2)
        number = p.add_run(str(note["number"]))
        number.font.superscript = True; number.font.size = Pt(7.5)
        number.font.color.rgb = RGBColor(0x7E, 0x18, 0x1E); _cjk(number)
        body = p.add_run("  " + inline_citations.format_source_note(note))
        body.font.size = Pt(8.5); body.font.color.rgb = RGBColor(0x60, 0x64, 0x70); _cjk(body)


def _h(doc, text: str, size: float = 13) -> None:
    from docx.shared import Pt, RGBColor
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x7E, 0x18, 0x1E)
    _cjk(r)


def _para(doc, text: str, size: float = 10.5, indent: float | None = None,
          gray: bool = False, bold: bool = False) -> None:
    from docx.shared import Pt, RGBColor
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.35
    if indent:
        p.paragraph_format.left_indent = Pt(indent)
    _inline(p, text, size=size, bold=bold,
            color=RGBColor(0x60, 0x64, 0x70) if gray else None)


def _dedupe_citations(msgs: list[dict]) -> dict[str, list[dict]]:
    """跨全会话去重引用，按类型分组，保持首次出现顺序。"""
    groups: dict[str, list[dict]] = {"statute": [], "case": [], "official": []}
    seen: set = set()
    for m in msgs:
        for c in (m.get("citations") or []):
            ctype = c.get("type") or ""
            if ctype == "statute":
                key = ("s", c.get("section_id") or (c.get("title"), c.get("section_label")))
                grp = "statute"
            elif ctype == "case":
                key = ("c", c.get("case_id") or c.get("document_id"))
                grp = "case"
            elif ctype in ("regulation", "guidance", "treaty", "amendment"):
                key = ("o", ctype, c.get("chunk_id") or c.get("title"))
                grp = "official"
            else:
                continue
            if key in seen:
                continue
            seen.add(key)
            groups[grp].append(c)
    return groups


_OFFICIAL_KIND = {"regulation": "官方公告/规章", "guidance": "监管指南",
                  "treaty": "条约/协定", "amendment": "法律沿革"}


def build_docx(conv: dict, msgs: list[dict], meta: dict | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

    today = date.today()
    ref_no = f"IDFL-{today.strftime('%Y%m%d')}-{conv['id']:04d}"

    # ---- 标题区 ----
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("泰国域外法查明意见书"); r.bold = True; r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x7E, 0x18, 0x1E); _cjk(r)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Thailand Foreign-Law Research Opinion"); r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x60, 0x64, 0x70); _cjk(r)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(f"编号：{ref_no}　　出具日期：{today.strftime('%Y 年 %m 月 %d 日')}　　出具平台：{config.PLATFORM_NAME}")
    r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x60, 0x64, 0x70); _cjk(r)
    # V1.3-B4.5：委托方 / 承办人 / 委托事项编号（可选，实务文书署名要素）
    meta = meta or {}
    meta_line = "　　".join(
        f"{lbl}：{meta[k]}" for lbl, k in
        [("委托方", "client_name"), ("承办人", "handler_name"), ("委托事项编号", "matter_ref")]
        if meta.get(k))
    if meta_line:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(meta_line)
        r.font.size = Pt(10); _cjk(r)

    # 分隔线（细边框段落）
    p = doc.add_paragraph()
    from docx.oxml.ns import qn as _qn
    from docx.oxml import OxmlElement
    pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(_qn("w:val"), "single"); bottom.set(_qn("w:sz"), "8")
    bottom.set(_qn("w:color"), "7E181E"); pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)

    questions = [m for m in msgs if m["role"] == "user"]
    answers = [m for m in msgs if m["role"] == "assistant"]
    groups = _dedupe_citations(msgs)

    # ---- 一、查明事项 ----
    _h(doc, "一、查明事项")
    if len(questions) == 1:
        _para(doc, questions[0]["content"])
    else:
        for i, q in enumerate(questions, 1):
            _para(doc, f"{i}. {q['content']}")

    # ---- 二、查明结论 ----
    _h(doc, "二、查明结论")
    # 按顺序配对：每条 assistant 前最近的 user 是它的问题
    qa_pairs = []
    pending_q = None
    for m in msgs:
        if m["role"] == "user":
            pending_q = m
        elif m["role"] == "assistant":
            qa_pairs.append((pending_q, m))
            pending_q = None
    note_offset = 0
    citation_history: list[dict] = []
    for i, (q, a) in enumerate(qa_pairs, 1):
        if len(qa_pairs) > 1:
            _para(doc, f"（{i}）针对查明事项 {i}：", bold=True, size=11)
        citation_history.extend(a.get("citations") or [])
        enriched = inline_citations.enrich_verification(
            a["content"], citation_history[-400:], a.get("verification") or {})
        anchors = enriched.get("inline_citations") or []
        local_max = max((int(anchor.get("number") or 0) for anchor in anchors), default=0)
        if note_offset:
            for anchor in anchors:
                anchor["number"] = int(anchor["number"]) + note_offset
        _md_to_docx(doc, a["content"], anchors)
        _add_inline_notes(doc, anchors)
        note_offset += local_max

    # ---- 三、成文法依据 ----
    if groups["statute"]:
        _h(doc, "三、成文法依据")
        for i, c in enumerate(groups["statute"], 1):
            head = f"{i}. **{c.get('title') or '——'}**"
            if c.get("act_number"):
                head += f"（{c['act_number']}）"
            if c.get("section_label"):
                head += f"，{_statute_locator(c)}"
            _para(doc, head, size=10.5)
            _para(
                doc,
                f"文本属性：{_statute_text_attribute(c)}",
                size=9,
                indent=14,
                gray=True,
            )
            text = (c.get("section_text") or "").strip()
            if text:
                excerpt = " ".join(text.split())[:500]
                _para(doc, f"条文摘录：{excerpt}{'…' if len(text) > 500 else ''}",
                      size=9.5, indent=14, gray=True)
            if c.get("source_url") or c.get("file_url"):
                _para(doc, f"官方出处：{c.get('source_url') or c.get('file_url')}",
                      size=9, indent=14, gray=True)

    # ---- 四、判例支撑 ----
    if groups["case"]:
        _h(doc, "四、判例支撑")
        for i, c in enumerate(groups["case"], 1):
            _para(doc, f"{i}. **{c.get('case_name') or '——'}**", size=10.5)
            meta = "；".join(x for x in [
                c.get("court"),
                f"法院级别：{c['court_level']}" if c.get("court_level") else None,
                f"判决日期：{c['decision_date']}" if c.get("decision_date") else None,
                f"案号：{c['case_number']}" if c.get("case_number") else None,
                f"引证：{c['citation']}" if c.get("citation") else None,
                f"库内判例 ID：{c['case_id']}" if c.get("case_id") else None,
            ] if x)
            if meta:
                _para(doc, meta, size=9.5, indent=14, gray=True)
            if c.get("source_url") or c.get("source_urn"):
                _para(doc, f"官方出处：{c.get('source_url') or c.get('source_urn')}",
                      size=9, indent=14, gray=True)

    # ---- 五、官方材料 ----
    if groups["official"]:
        _h(doc, "五、官方公告、监管指南与条约")
        for i, c in enumerate(groups["official"], 1):
            kind = _OFFICIAL_KIND.get(c.get("type"), "官方材料")
            head = f"{i}. **{c.get('title') or '——'}**（{kind}"
            for extra in (c.get("document_type"),
                          f"Gazette {c['gazette_number']}" if c.get("gazette_number") else None,
                          c.get("source_authority"), c.get("publication_date")):
                if extra:
                    head += f"，{extra}"
            head += "）"
            _para(doc, head, size=10.5)
            if c.get("source_url") or c.get("file_url"):
                _para(doc, f"官方出处：{c.get('source_url') or c.get('file_url')}",
                      size=9, indent=14, gray=True)

    # ---- 六、查明方法与局限 ----
    _h(doc, "六、查明方法与局限说明")
    n_search = 0
    tools_used: set[str] = set()
    for a in answers:
        for t in (a.get("tool_trace") or []):
            if t.get("tool"):
                tools_used.add(t["tool"])
                n_search += 1
    _para(doc,
          f"本意见书由{config.PLATFORM_NAME}生成。平台以大语言模型智能体为核心，"
          f"对example_database（OCS 官方泰文法规目录与重点法律英文译本；"
          f"英文材料均以泰语官方文本为准）"
          f"执行「PostgreSQL 全文检索 + 向量语义检索 + 交叉编码重排」的三路混合检索，"
          f"并遵循「先取证、后作答」的约束：结论仅基于检索到的真实材料综合形成，全部依据均附出处，"
          f"正文中的成文法引用另经规则引擎逐条回查库内验真。")
    if n_search:
        _para(doc, f"本次查明共执行 {n_search} 次检索取证，涉及工具：{('、'.join(sorted(tools_used)))}。",
              size=9.5, gray=True)
    # 引用核验汇总（V1.2 反臆造；V1.3 修复：补 declared 档，四档之和=总数）
    v_sum = {"verified": 0, "partial": 0, "declared": 0, "unverified": 0, "total": 0}
    for a in answers:
        s = ((a.get("verification") or {}).get("summary")) or {}
        for k in v_sum:
            v_sum[k] += int(s.get(k) or 0)
    if v_sum["total"]:
        _para(doc,
              f"机器核验：正文中的成文法引用已经规则引擎自动验真（共 {v_sum['total']} 项）——"
              f"已核验 {v_sum['verified']} 项（见于本轮检索材料原文）、"
              f"部分核验 {v_sum['partial']} 项（库内确认存在但未经本轮取证）、"
              f"已声明 {v_sum['declared']} 项（正文已如实注明未经库内核验）、"
              f"未核验 {v_sum['unverified']} 项（建议人工复核）。", size=9.5, gray=True)
    _para(doc,
          "局限说明：本意见书基于平台库内材料的检索结果自动汇编，供研究与实务参考；"
          "库内材料存在时点与覆盖范围限制，正文中已如实标注未能检索到的事项。"
          "引用与援引请以官方原始文本为准；本意见书不构成正式法律意见。", size=9.5, gray=True)

    # ---- 附录：来源清单 ----
    all_srcs = groups["statute"] + groups["case"] + groups["official"]
    if all_srcs:
        _h(doc, "附录　来源清单")
        t = doc.add_table(rows=1 + len(all_srcs), cols=3)
        t.style = "Table Grid"
        for ci, htxt in enumerate(["类型", "名称", "官方出处"]):
            cell = t.cell(0, ci); cell.paragraphs[0].text = ""
            _inline(cell.paragraphs[0], htxt, size=9, bold=True)
        type_label = {"statute": "成文法", "case": "判例"}
        for ri, c in enumerate(all_srcs, 1):
            label = type_label.get(c.get("type")) or _OFFICIAL_KIND.get(c.get("type"), "材料")
            name = c.get("title") or c.get("case_name") or "——"
            if c.get("section_label"):
                name += f" {_statute_locator(c)}"
            url = c.get("source_url") or c.get("file_url") or c.get("source_urn") or "库内存档"
            for ci, v in enumerate([label, name, url]):
                cell = t.cell(ri, ci); cell.paragraphs[0].text = ""
                _inline(cell.paragraphs[0], str(v)[:200], size=8.5)

    # ---- 落款 ----
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _AL
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(22)
    p.alignment = _AL.RIGHT
    r = p.add_run(config.PLATFORM_NAME); r.font.size = Pt(11); r.bold = True; _cjk(r)
    p = doc.add_paragraph(); p.alignment = _AL.RIGHT
    r = p.add_run(today.strftime("%Y 年 %m 月 %d 日")); r.font.size = Pt(10); _cjk(r)

    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_to_pdf(docx_bytes: bytes) -> bytes:
    """LibreOffice headless 转 PDF。独立 profile + HOME=/tmp，zafl 可运行；串行锁防并发损坏。"""
    with _soffice_lock:
        with tempfile.TemporaryDirectory(prefix="my_report_") as td:
            src = os.path.join(td, "report.docx")
            with open(src, "wb") as f:
                f.write(docx_bytes)
            # profile 按 uid 隔离：dev(root) 与 prod(zafl) 不共享，避免属主冲突
            env = dict(os.environ, HOME="/tmp")
            proc = subprocess.run(
                [SOFFICE, "--headless",
                 f"-env:UserInstallation=file:///tmp/my_lo_profile_{os.getuid()}",
                 "--convert-to", "pdf", "--outdir", td, src],
                capture_output=True, timeout=120, env=env,
            )
            pdf = os.path.join(td, "report.pdf")
            if not os.path.exists(pdf):
                logger.error("soffice convert failed: %s", proc.stderr.decode(errors="ignore")[:500])
                raise RuntimeError("PDF 转换失败")
            with open(pdf, "rb") as f:
                return f.read()


# ---------------- 路由 ----------------

@router.get("/{conv_id}/report")
def export_report(conv_id: int, format: str = "pdf",
                  client_name: str | None = None, handler_name: str | None = None,
                  matter_ref: str | None = None,
                  user: dict = Depends(deps.get_current_user)):
    if format not in ("pdf", "docx"):
        raise HTTPException(status_code=400, detail="format 仅支持 pdf / docx")
    # V1.3-B4.6 单用户导出限额：每日 30 份（audit 计数），防批量拖库
    n_today = db.query_one(
        """SELECT count(*) AS n FROM app_audit_log
           WHERE user_id = %s AND action = 'export_report' AND created_at > now() - interval '1 day'""",
        [user["id"]])
    if n_today and n_today["n"] >= 30:
        raise HTTPException(status_code=429, detail="今日导出次数已达上限（30 份），请明日再试")
    meta = {"client_name": (client_name or "")[:60] or None,
            "handler_name": (handler_name or "")[:60] or None,
            "matter_ref": (matter_ref or "")[:60] or None}
    conv = db.query_one(
        "SELECT * FROM app_conversations WHERE id = %s AND user_id = %s",
        [conv_id, user["id"]],
    )
    if not conv:
        raise HTTPException(status_code=404, detail="会话不存在")
    msgs = db.query_all(
        """
        SELECT role, content, citations, tool_trace, verification
        FROM app_messages WHERE conversation_id = %s ORDER BY created_at, id
        """,
        [conv_id],
    )
    if not any(m["role"] == "assistant" and (m.get("content") or "").strip() for m in msgs):
        raise HTTPException(status_code=400, detail="该会话尚无查明结论，无法导出")

    if not _export_slots.acquire(blocking=False):
        raise HTTPException(status_code=429, detail="导出请求繁忙，请稍后再试")
    try:
        docx_bytes = build_docx(conv, msgs, meta)
        if format == "docx":
            content, mt, ext = docx_bytes, \
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
        else:
            content, mt, ext = _docx_to_pdf(docx_bytes), "application/pdf", "pdf"
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("report export failed (conv=%s): %s", conv_id, exc)
        raise HTTPException(status_code=500, detail="意见书生成失败，请稍后重试")
    finally:
        _export_slots.release()

    db.audit(user["id"], "export_report", {"conversation_id": conv_id, "format": format})
    fname = f"泰国域外法查明意见书_{date.today().strftime('%Y%m%d')}_{conv_id:04d}.{ext}"
    return Response(
        content=content, media_type=mt,
        headers={"Content-Disposition":
                 f"attachment; filename=\"report.{ext}\"; filename*=UTF-8''{quote(fname)}"},
    )
